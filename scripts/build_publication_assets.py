"""Build portfolio PDFs, a 12-slide deck, an Excel dictionary, and SVG diagrams."""

from __future__ import annotations

import argparse
import html
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]

DATA_DICTIONARY = {
    "Sales Transactions": [
        ("order_id", "string", "Unique transaction/day identifier."),
        ("order_date", "date", "Transaction date; parseable by pandas."),
        ("sku_id", "string", "Product identifier matching the product table."),
        ("quantity", "integer", "Units; negative rows represent returns."),
        ("unit_price", "decimal", "Non-negative selling-price value."),
        ("discount", "decimal", "Discount fraction from 0 through 1."),
        ("promotion", "boolean", "Optional descriptive promotion flag."),
    ],
    "Product Assumptions": [
        ("sku_id", "string", "Unique SKU identifier."),
        ("product_name", "string", "Product label."),
        ("category", "string", "Merchandising category."),
        ("unit_cost", "decimal", "Non-negative product cost; ABC proxy input."),
        ("current_stock", "number", "On-hand stock snapshot at planning date."),
        ("lead_time_days", "number", "Non-negative assumed lead time in days."),
        ("minimum_order_quantity", "number", "Supplier minimum order amount."),
        ("order_cost", "decimal", "Fixed cost per replenishment order."),
        ("annual_holding_cost", "decimal", "Annual holding cost per unit; must be positive."),
    ],
    "Derived Outputs": [
        ("date", "date", "SKU-day historical or forecast date."),
        ("demand", "number", "Non-negative net daily demand after return netting."),
        ("forecast", "number", "Non-negative predicted SKU-day demand."),
        ("mae / rmse / wape", "number", "Chronological holdout error metrics."),
        ("safety_stock / reorder_point", "number", "Policy estimates in units."),
        ("recommended_order", "integer", "Review-only replenishment suggestion."),
        ("abc_xyz", "string", "Value/variability segment."),
    ],
    "Assumptions": [
        ("Data provenance", "synthetic", "Bundled sample data and supplier assumptions are illustrative."),
        ("Forecast split", "chronological", "Final contiguous observations are held out per SKU."),
        ("Safety stock", "approximation", "Z x daily demand standard deviation x sqrt(fixed lead time)."),
        ("Decision status", "human review", "Recommendations are not purchase orders or business guarantees."),
    ],
}


def ascii_text(value: str) -> str:
    return unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")


def markdown_story(path: Path, styles: dict[str, ParagraphStyle]) -> list[object]:
    story: list[object] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        content = line.strip()
        if not content or content.startswith("```"):
            if content.startswith("```"):
                story.append(Spacer(1, 3))
            continue
        if content.startswith("#"):
            level = len(content) - len(content.lstrip("#"))
            style = styles["Title"] if level == 1 else styles["Heading2"] if level > 2 else styles["Heading1"]
            story.append(Paragraph(html.escape(ascii_text(content.lstrip("# "))), style))
            story.append(Spacer(1, 3))
            continue
        content = re.sub(r"`([^`]+)`", r"\1", content)
        content = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", content)
        if content.startswith(">"):
            content = f"<i>{html.escape(ascii_text(content.lstrip('> ')))}</i>"
        elif content.startswith(("- ", "* ")):
            content = "&#8226; " + html.escape(ascii_text(content[2:]))
        else:
            content = html.escape(ascii_text(content))
        story.append(Paragraph(content, styles["BodyText"]))
        story.append(Spacer(1, 4))
    return story


def build_pdf(path: Path, title: str, body: list[object]) -> None:
    styles = getSampleStyleSheet()
    styles["Title"].alignment = TA_CENTER
    styles["Title"].textColor = colors.HexColor("#12304A")
    styles["Heading1"].textColor = colors.HexColor("#1E5B78")
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm, title=title,
        author="Demand Forecasting & Inventory Optimization project",
    )
    doc.build(body, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)


def _pdf_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#637381"))
    canvas.drawString(18 * mm, 10 * mm, "Synthetic portfolio demonstration - verify assumptions before use")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))
    canvas.restoreState()


def project_report_story(results: dict) -> list[object]:
    styles = getSampleStyleSheet()
    styles["Title"].alignment = TA_CENTER
    overall = results["metrics"].query("sku_id == 'ALL'").sort_values("wape")
    winner = overall.iloc[0]
    reorder_count = int((results["decisions"]["recommended_order"] > 0).sum())
    rows = [["Model", "MAE (units)", "RMSE (units)", "WAPE"]]
    for row in overall.itertuples():
        rows.append([
            str(row.model), f"{row.mae:.2f}", f"{row.rmse:.2f}",
            "N/A" if row.wape is None else f"{row.wape:.1%}",
        ])
    table = Table(rows, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CAD5DF")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F6FA")]),
        ("PADDING", (0, 0), (-1, -1), 7),
    ]))
    return [
        Paragraph("Project Report", styles["Title"]),
        Paragraph(html.escape("Demand Forecasting & Inventory Optimization"), styles["Heading1"]),
        Paragraph(
            "This report is generated from the repository's deterministic synthetic sample. "
            "It demonstrates method execution only; results are not business evidence.",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("Scope and workflow", styles["Heading2"]),
        Paragraph(
            f"The sample contains {results['products']['sku_id'].nunique()} SKUs and "
            f"{results['demand']['date'].nunique()} daily dates. Transactions are cleaned, "
            "net returns into SKU-day demand, and are evaluated with a chronological holdout. "
            "The selected forecast is translated into service-level stock policy estimates.",
            styles["BodyText"],
        ),
        Spacer(1, 8),
        Paragraph("Illustrative holdout comparison", styles["Heading2"]),
        table,
        Spacer(1, 8),
        Paragraph(
            f"Best WAPE on this synthetic holdout: {winner.model} ({winner.wape:.1%}). "
            f"Suggested replenishment review count: {reorder_count} of {len(results['decisions'])} SKUs. "
            "Neither value should be interpreted as an expected result on other data.",
            styles["BodyText"],
        ),
        Spacer(1, 8),
        Paragraph("Decision formulas", styles["Heading2"]),
        Paragraph(
            "Safety stock = Z x daily demand standard deviation x sqrt(fixed lead time). "
            "Reorder point = mean daily demand x lead time + safety stock. "
            "EOQ = sqrt(2 x annual demand x order cost / annual holding cost per unit). "
            "These are simplified assumptions and omit purchase orders, uncertain lead time, "
            "capacity, expiry, and supplier calendars.",
            styles["BodyText"],
        ),
        Spacer(1, 8),
        Paragraph("Recommended interpretation", styles["Heading2"]),
        Paragraph(
            "Use this project to demonstrate reproducible engineering, time-series validation, "
            "and explainable policies. Do not claim inventory savings, recovered demand, or "
            "guaranteed service levels without a validated operational pilot.",
            styles["BodyText"],
        ),
    ]


def build_dictionary(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, records in DATA_DICTIONARY.items():
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Field", "Type", "Definition"])
        for row in records:
            sheet.append(list(row))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12304A")
        for column, width in {"A": 30, "B": 20, "C": 95}.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def build_deck(path: Path, results: dict) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    metrics = results["metrics"].query("sku_id == 'ALL'").sort_values("wape")
    best = metrics.iloc[0]
    decision_count = int((results["decisions"]["recommended_order"] > 0).sum())
    slides = [
        ("Demand Forecasting & Inventory Optimization", ["End-to-end portfolio demonstration", "Synthetic data; validate before operational use"]),
        ("Business problem", ["Demand uncertainty and supplier lead time complicate replenishment.", "Goal: make demand outlook and stock-policy assumptions reviewable."]),
        ("Decision questions", ["What demand is plausible?", "How does the forecast compare with baselines?", "Which SKUs merit stock review?"]),
        ("End-to-end workflow", ["Transactions -> quality checks -> SKU-day demand -> features", "Chronological evaluation -> forecast -> inventory policy -> human review"]),
        ("Data and provenance", [f"{results['products']['sku_id'].nunique()} fictional SKUs and {results['demand']['date'].nunique()} sample days.", "All sample transactions and supplier/stock inputs are synthetic."]),
        ("Data quality and EDA", ["Deduplicate transaction IDs; validate dates, price, discount and SKU.", "Preserve returns, net by day, then floor demand at zero.", "Review demand variability and no-sale-day frequency."]),
        ("Features and leakage control", ["Lag and rolling demand features are shifted to prior dates.", "Calendar fields are known in advance.", "Future lag values are generated recursively."]),
        ("Forecast models", ["Last-observation naive", "Seven-day seasonal naive", "Seven-day moving average", "Histogram gradient boosting"]),
        ("Synthetic holdout results", [f"Chronological holdout best WAPE: {best.wape:.1%} ({best.model}).", f"MAE: {best.mae:.2f}; RMSE: {best.rmse:.2f} units.", "Illustrative only; no claim of generalized accuracy."]),
        ("Inventory policy", ["Safety stock: Z x demand standard deviation x sqrt(lead time).", "Reorder point: expected lead-time demand + safety stock.", "EOQ and supplier MOQ are assumptions to validate."]),
        ("Application and deliverables", ["Streamlit dashboard and CSV upload mode.", "FastAPI read-only demo endpoint.", "SQL examples, notebooks, CI tests, Docker and docs."]),
        ("Limitations and next steps", [f"{decision_count} of {len(results['decisions'])} demo SKUs are flagged to review.", "No open POs, live ERP, uncertain lead-time model, or autonomous ordering.", "Pilot only with validated data, monitoring, and human approvals."]),
    ]
    for index, (heading, bullets) in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(246, 249, 252)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.text = heading
        title_frame.paragraphs[0].font.size = Pt(30)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].font.color.rgb = RGBColor(18, 48, 74)
        body = slide.shapes.add_textbox(Inches(0.9), Inches(1.65), Inches(11.5), Inches(4.9))
        frame = body.text_frame
        frame.word_wrap = True
        for item_index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if item_index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(22)
            paragraph.font.color.rgb = RGBColor(44, 62, 80)
            paragraph.space_after = Pt(19)
        footer = slide.shapes.add_textbox(Inches(0.7), Inches(7.05), Inches(12), Inches(0.25))
        footer.text_frame.text = f"Synthetic portfolio demonstration  |  {index + 1:02d} / {len(slides):02d}"
        footer.text_frame.paragraphs[0].font.size = Pt(10)
        footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(99, 115, 129)
    deck.save(path)


def build_svg_diagrams(output_dir: Path) -> None:
    stages = [
        ("Data sources", "Sales + SKU assumptions"),
        ("Data engineering", "Validate, clean, aggregate"),
        ("Forecasting", "Features + time split"),
        ("Inventory", "Safety stock + reorder"),
        ("Decision support", "Dashboard + human review"),
    ]
    blocks = []
    for index, (title, description) in enumerate(stages):
        x = 35 + index * 230
        blocks.append(
            f'<rect x="{x}" y="60" width="190" height="105" rx="12" fill="#F1F6FA" stroke="#1E5B78" stroke-width="2"/>'
            f'<text x="{x + 95}" y="103" text-anchor="middle" font-size="18" font-family="Arial" fill="#12304A">{title}</text>'
            f'<text x="{x + 95}" y="133" text-anchor="middle" font-size="12" font-family="Arial" fill="#445">{description}</text>'
        )
        if index < len(stages) - 1:
            blocks.append(f'<path d="M {x + 195} 112 L {x + 222} 112" stroke="#526D82" stroke-width="3" marker-end="url(#arrow)"/>')
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" width="1220" height="230" viewBox="0 0 1220 230">'
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto">'
        '<path d="M0,0 L0,6 L9,3 z" fill="#526D82"/></marker></defs>'
        '<rect width="100%" height="100%" fill="white"/><text x="35" y="35" font-size="16" font-family="Arial" fill="#526D82">'
        'Synthetic-data reference architecture</text>' + "".join(blocks) + "</svg>"
    )
    (output_dir / "system_architecture.svg").write_text(svg, encoding="utf-8")
    (output_dir / "inventory_workflow.svg").write_text(svg.replace(
        "Synthetic-data reference architecture", "Inventory decision workflow"
    ), encoding="utf-8")
    (output_dir / "forecasting_workflow.svg").write_text(svg.replace(
        "Synthetic-data reference architecture", "Forecasting and deployment workflow"
    ), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "reports")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    from demand_inventory.service import run_project

    results = run_project(horizon=30, test_days=28)
    markdown_docs = [
        ("Complete_Project_Explanation.pdf", "Complete Project Explanation", ROOT / "docs/project-explanation.md"),
        ("Project_Workflow.pdf", "Project Workflow", ROOT / "docs/workflow.md"),
        ("Interview_Questions_Answers.pdf", "Interview Questions and Answers", ROOT / "docs/interview-guide.md"),
        ("Setup_Guide.pdf", "Setup Guide", ROOT / "README.md"),
        ("Methodology_Guide.pdf", "Methodology Guide", ROOT / "docs/methodology.md"),
    ]
    styles = getSampleStyleSheet()
    for filename, title, source in markdown_docs:
        build_pdf(args.output_dir / filename, title, markdown_story(source, styles))
    build_pdf(
        args.output_dir / "Project_Report.pdf", "Project Report",
        project_report_story(results),
    )
    build_dictionary(args.output_dir / "data_dictionary.xlsx")
    build_deck(args.output_dir / "Project_Presentation.pptx", results)
    build_svg_diagrams(args.output_dir)
    print(f"Publication package written to {args.output_dir}")
    print("Generated: 6 PDFs, 1 XLSX data dictionary, 1 12-slide PPTX, 3 SVG diagrams.")


if __name__ == "__main__":
    main()
